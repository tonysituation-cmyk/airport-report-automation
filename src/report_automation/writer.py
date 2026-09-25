"""Render aggregated tables into a formatted, charted Excel workbook
with openpyxl — the artefact managers actually open."""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .alerts import Alert

HEADER_FILL = PatternFill("solid", fgColor="1A2E4A")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="1A2E4A")
ALERT_TITLE_FONT = Font(bold=True, size=12, color="B03030")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")


def _write_table(ws: Worksheet, df: pd.DataFrame, start_row: int = 1,
                 start_col: int = 1) -> int:
    """Write a frame as a styled table; returns the last data row."""
    for j, col in enumerate(df.columns, start=start_col):
        cell = ws.cell(row=start_row, column=j, value=col.replace("_", " ").title())
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    is_date: dict[int, bool] = {}
    for i, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for j, col in enumerate(df.columns, start=start_col):
            value: Any = row[col]
            if hasattr(value, "to_pydatetime"):
                value = value.to_pydatetime()
                is_date[j] = True
            cell = ws.cell(row=i, column=j, value=value)
            if is_date.get(j):
                cell.number_format = "yyyy-mm-dd"
    for j, col in enumerate(df.columns, start=start_col):
        if is_date.get(j):
            width = 12  # fits yyyy-mm-dd
        else:
            longest = max((len(str(v)) for v in df[col]), default=8)
            width = max(len(str(col)) + 4, longest + 3, 12)
        ws.column_dimensions[get_column_letter(j)].width = width
    return start_row + len(df)


def write_report(
    summary: pd.DataFrame,
    events: pd.DataFrame,
    alerts: list[Alert],
    path: Path,
    *,
    delay_rate_threshold: float,
) -> Path:
    """Write the full workbook (summary + charts, events, alerts) to ``path``."""
    wb = Workbook()

    # --- Daily Summary sheet ---
    ws = wb.active
    ws.title = "Daily Summary"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Weekly Operations Report"
    ws["A1"].font = TITLE_FONT
    last = _write_table(ws, summary, start_row=3)
    ws.freeze_panes = "A4"

    col_index = {name: list(summary.columns).index(name) + 1 for name in summary.columns}
    delay_col_letter = get_column_letter(col_index["delay_rate"])
    ws.conditional_formatting.add(
        f"{delay_col_letter}4:{delay_col_letter}{last}",
        CellIsRule(operator="greaterThan",
                   formula=[str(delay_rate_threshold)], fill=RED_FILL),
    )
    ws.auto_filter.ref = f"A3:{get_column_letter(len(summary.columns))}{last}"

    categories = Reference(ws, min_col=col_index["date"], min_row=4, max_row=last)

    bar = BarChart()
    bar.title = "Passengers per Day"
    bar.y_axis.title = "Passengers"
    bar.add_data(Reference(ws, min_col=col_index["passengers"], min_row=3, max_row=last),
                 titles_from_data=True)
    bar.set_categories(categories)
    bar.height, bar.width = 8, 16
    ws.add_chart(bar, "J3")

    line = LineChart()
    line.title = "Delay Rate"
    line.y_axis.title = "Delay Rate"
    line.add_data(Reference(ws, min_col=col_index["delay_rate"], min_row=3, max_row=last),
                  titles_from_data=True)
    line.set_categories(categories)
    line.height, line.width = 8, 16
    ws.add_chart(line, "J20")

    # --- Ground Events sheet ---
    ws2 = wb.create_sheet("Ground Events")
    ws2.sheet_view.showGridLines = False
    _write_table(ws2, events)

    # --- Alerts sheet ---
    ws3 = wb.create_sheet("Alerts")
    ws3.sheet_view.showGridLines = False
    ws3["A1"] = "Threshold Alerts"
    ws3["A1"].font = ALERT_TITLE_FONT
    if alerts:
        for i, alert in enumerate(alerts, start=3):
            ws3.cell(row=i, column=1,
                     value=f"{alert.date} | {alert.rule} | {alert.message}")
    else:
        ws3.cell(row=3, column=1, value="No threshold breaches this period.")
    ws3.column_dimensions["A"].width = 80

    wb.properties.title = "Weekly Operations Report"
    wb.properties.creator = "airport-report-automation"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
