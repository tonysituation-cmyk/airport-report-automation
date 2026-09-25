from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from report_automation.alerts import Alert
from report_automation.writer import write_report


def _summary() -> pd.DataFrame:
    return pd.DataFrame({
        "date": pd.to_datetime(["2026-09-07", "2026-09-08"]),
        "flights": [24, 26],
        "delayed_flights": [1, 0],
        "avg_delay_min": [3.0, 0.0],
        "delay_rate": [0.042, 0.0],
        "passengers": [9709, 9830],
    })


def _events() -> pd.DataFrame:
    return pd.DataFrame({
        "date": pd.to_datetime(["2026-09-07"]),
        "gate_change": [3],
        "crew_delay": [1],
    })


def test_write_report_produces_workbook_with_charts(tmp_path):
    out = tmp_path / "report.xlsx"
    write_report(_summary(), _events(),
                 [Alert(date=datetime(2026, 9, 8).date(), rule="delay_rate",
                        message="breach")],
                 out, delay_rate_threshold=0.20)
    assert out.exists() and out.stat().st_size > 5000

    wb = load_workbook(out)
    assert wb.sheetnames == ["Daily Summary", "Ground Events", "Alerts"]

    ws = wb["Daily Summary"]
    assert len(ws._charts) == 2
    assert ws["A1"].value == "Weekly Operations Report"

    alerts_ws = wb["Alerts"]
    assert "breach" in str(alerts_ws["A3"].value)


def test_write_report_marks_clean_run(tmp_path):
    out = tmp_path / "report.xlsx"
    write_report(_summary(), _events(), [], out, delay_rate_threshold=0.20)
    wb = load_workbook(out)
    assert "No threshold breaches" in str(wb["Alerts"]["A3"].value)
