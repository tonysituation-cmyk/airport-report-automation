import logging

import pytest
from openpyxl import load_workbook

from report_automation.pipeline import run


def test_end_to_end_run_produces_report(sandbox, caplog):
    with caplog.at_level(logging.INFO, logger="report_automation"):
        output = run(sandbox)

    assert output.exists() and output.stat().st_size > 5000
    assert "Run finished" in caplog.text
    assert any("ALERT" in rec.message for rec in caplog.records), \
        "sample data should trigger at least one threshold alert"

    wb = load_workbook(output)
    assert "Daily Summary" in wb.sheetnames


def test_run_fails_cleanly_on_missing_source(sandbox):
    sandbox.sqlite_db.unlink()  # simulate missing ops DB
    from report_automation.exceptions import DataSourceError

    with pytest.raises(DataSourceError, match="Ops database not found"):
        run(sandbox)
    # failure notification must be logged
    assert (sandbox.logs_dir / "alerts.log").exists()
